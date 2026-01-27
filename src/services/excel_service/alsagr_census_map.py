import openpyxl 
import pandas
import os
from src.utils.load_yaml import ATTACHMENTS_SAVE_DIR,ALSAGR_TEMPLATES_DIR,REFERRAL_FILE_STORE_DIR, ALSAGR_GENERATED_CENSUS_DIR
from src.utils.support_functions import get_replaced_referral_id
from datetime import datetime
from src.utils.logger import logger


def alsagr_map_census_data(id):
    census_filename = ""
    if id == 'default':
        for file_name in os.listdir(ATTACHMENTS_SAVE_DIR):
            if not file_name.startswith("Medical__") and file_name.endswith(".xlsx"):
                census_filename = file_name

        excel_data_df = pandas.read_excel(os.path.join(ATTACHMENTS_SAVE_DIR, census_filename), sheet_name='Sheet1')
        nationality_df = pandas.read_excel(os.path.join(ATTACHMENTS_SAVE_DIR, census_filename), sheet_name='Nationality_Updated')

    else:
        for file_name in os.listdir(os.path.join(REFERRAL_FILE_STORE_DIR, get_replaced_referral_id(id))):
            if not file_name.startswith("Medical__") and file_name.endswith(".xlsx"):
                census_filename = file_name

        excel_data_df = pandas.read_excel(os.path.join(REFERRAL_FILE_STORE_DIR, get_replaced_referral_id(id), census_filename), sheet_name='Sheet1')
        nationality_df = pandas.read_excel(os.path.join(REFERRAL_FILE_STORE_DIR, get_replaced_referral_id(id), census_filename), sheet_name='National_Updated')

    merged_df = pandas.merge(excel_data_df, nationality_df, left_on='Nationality', right_on='AL SAGR', how='left')

    wb = openpyxl.load_workbook(os.path.join(ALSAGR_TEMPLATES_DIR, "MemberUpload.xlsx"))
    ws = wb['Sheet1']

    for index, row in merged_df.iterrows():
        dob = row['DOB']
        # logger.info(f"DOB: {dob}")
        
        if isinstance(dob, str):
            dob = dob.strip()
            formats = ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"]  # Possible formats

            for fmt in formats:
                try:
                    dob = datetime.strptime(dob, fmt)
                    # logger.info(f"Parsed DOB: {dob}")
                    break  # Stop checking after first successful parse
                except ValueError:
                    continue
            else:
                dob = None  # No valid format found

        elif pandas.isna(dob):
            dob = None  # Handle NaN values

        if dob:
            dob_cell = ws.cell(row=index+2, column=2)
            dob_cell.value = dob
            # logger.info(f"Final DOB written: {dob}")
            dob_cell.number_format = "D-MM-YY"
        else:
            ws.cell(row=index+2, column=2).value = "Invalid DOB"
            logger.info(f"Invalid DOB: {dob}")
            
        ws.cell(row=index+2, column=1).value = row['Beneficiary First Name']
        ws.cell(row=index+2, column=3).value = row['Gender']
        ws.cell(row=index+2, column=4).value = row['Category']
        ws.cell(row=index+2, column=5).value = row['Nationality']
        ws.cell(row=index+2, column=6).value = row['Relation']
        ws.cell(row=index+2, column=7).value = row['Marital status']
        ws.cell(row=index+2, column=8).value = row['Visa Issued Emirates']
        ws.cell(row=index+2, column=9).value = row['Monthly salary ']
        ws.cell(row=index+2, column=10).value = row['Salary Type']
        ws.cell(row=index+2, column=11).value = row['Status']

    wb.save(os.path.join(ALSAGR_GENERATED_CENSUS_DIR, "MemberUpload.xlsx"))
